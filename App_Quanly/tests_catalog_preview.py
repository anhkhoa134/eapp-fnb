import io
from decimal import Decimal
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from App_Accounts.models import User
from App_Catalog.models import Category, Product, ProductUnit, StoreProduct, Topping, ProductTopping
from App_Core.models import AuditLog
from App_Tenant.models import Store, Tenant, UserStoreAccess


class CatalogPreviewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(name='Preview', public_slug='preview')
        cls.store = Store.objects.create(tenant=cls.tenant, name='Store', is_default=True)
        cls.manager = User.objects.create_user(username='preview_manager', tenant=cls.tenant, role=User.Role.MANAGER)
        UserStoreAccess.objects.create(user=cls.manager, store=cls.store, is_default=True)

    def setUp(self):
        self.client.force_login(self.manager)

    def workbook(self, rows=None, headers=None):
        wb = Workbook()
        wb.active.title = 'San_pham'
        wb.active.append(headers or ['ten_danh_muc', 'ten_san_pham', 'ten_don_vi', 'gia'])
        for row in rows if rows is not None else [['Nước', 'Trà', 'Ly', 20000]]:
            wb.active.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def preview(self, raw=None):
        return self.client.post(reverse('App_Quanly:catalog_import_preview'), {
            'excel_file': SimpleUploadedFile('test.xlsx', raw if raw is not None else self.workbook()),
        })

    def confirm(self, raw, token=''):
        return self.client.post(reverse('App_Quanly:catalog_import_upload'), {
            'excel_file': SimpleUploadedFile('test.xlsx', raw), 'preview_token': token,
        }, HTTP_ACCEPT='application/json')

    def existing(self, *, price=15000, active=True):
        category = Category.objects.create(tenant=self.tenant, name='Nước', is_active=active)
        product = Product.objects.create(tenant=self.tenant, category=category, name='Trà', description='Giữ mô tả', is_active=active)
        unit = ProductUnit.objects.create(product=product, name='Ly', price=price, is_active=active)
        return category, product, unit

    def test_scan_creates_no_catalog_or_audit_records(self):
        before = AuditLog.objects.count()
        response = self.preview().json()
        self.assertTrue(response['ok'], response['errors'])
        self.assertEqual(response['summary']['create'], 1)
        self.assertEqual(response['summary']['categories_new'], 1)
        self.assertEqual(response['summary']['products_new'], 1)
        self.assertEqual(response['rows'][0]['action'], 'create')
        self.assertFalse(Category.objects.exists())
        self.assertFalse(Product.objects.exists())
        self.assertEqual(AuditLog.objects.count(), before)

    def test_update_unchanged_and_new_unit_preview_matches_import(self):
        _, product, unit = self.existing()
        raw = self.workbook([['Nước', 'Trà', 'Ly', 20000], ['Nước', 'Trà', 'Chai', 25000]])
        result = self.preview(raw).json()
        self.assertEqual([row['action'] for row in result['rows']], ['update', 'create'])
        self.assertEqual(Decimal(result['rows'][0]['old_price']), Decimal('15000'))
        unit.refresh_from_db()
        self.assertEqual(unit.price, Decimal('15000'))
        self.assertTrue(self.confirm(raw, result['preview_token']).json()['ok'])
        unit.refresh_from_db()
        product.refresh_from_db()
        self.assertEqual(unit.price, Decimal('20000'))
        self.assertEqual(product.description, 'Giữ mô tả')
        self.assertEqual(product.units.count(), 2)
        result = self.preview(raw).json()
        self.assertEqual([row['action'] for row in result['rows']], ['unchanged', 'unchanged'])
        timestamp = unit.updated_at
        self.assertTrue(self.confirm(raw, result['preview_token']).json()['ok'])
        unit.refresh_from_db()
        self.assertEqual(unit.updated_at, timestamp)

    def test_same_product_many_units_counts_entities_once(self):
        result = self.preview(self.workbook([['Nước', 'Trà', 'Ly', 0], ['Nước', 'Trà', 'Chai', '25.000']])).json()
        self.assertTrue(result['ok'], result['errors'])
        self.assertEqual(result['summary']['categories_new'], 1)
        self.assertEqual(result['summary']['products_new'], 1)
        self.assertEqual(result['summary']['create'], 2)

    def test_inactive_custom_price_and_hidden_store_warnings(self):
        _, product, _ = self.existing(active=False)
        StoreProduct.objects.create(store=self.store, product=product, is_available=False, custom_price=25000)
        result = self.preview().json()
        self.assertTrue(result['ok'], result['errors'])
        warnings = ' '.join(result['rows'][0]['warnings'])
        self.assertIn('đang tắt', warnings)
        self.assertIn('giá riêng', warnings)
        self.assertIn('giới hạn', warnings)

    def test_same_name_in_another_category_creates_separate_product(self):
        self.existing()
        raw = self.workbook([['Món khác', 'Trà', 'Ly', 20000]])
        result = self.preview(raw).json()
        self.assertTrue(result['ok'], result['errors'])
        self.assertIn('danh mục khác', ' '.join(result['rows'][0]['warnings']))
        self.assertTrue(self.confirm(raw, result['preview_token']).json()['ok'])
        self.assertEqual(Product.objects.filter(name='Trà').count(), 2)

    def test_case_difference_warned_and_outer_whitespace_trimmed(self):
        self.existing()
        result = self.preview(self.workbook([[' Nước ', ' Trà ', ' Ly ', 15000]])).json()
        self.assertEqual(result['rows'][0]['action'], 'unchanged')
        result = self.preview(self.workbook([['nước', 'trà', 'Ly', 15000]])).json()
        self.assertTrue(result['ok'])
        self.assertIn('gần giống', ' '.join(result['rows'][0]['warnings']))

    def test_invalid_rows_block_confirmation_and_keep_valid_rows_visible(self):
        invalid = [
            ['Nước', 'Trà', 'Ly', 20000],  # Duplicate row.
            ['', 'Trà', 'Ly', 10000], ['Nước', '', 'Ly', 10000], ['Nước', 'Trà', '', 10000],
            ['Nước', 'Khác', 'Ly', -1], ['Nước', 'Khác', 'Ly', 'abc'],
            ['Nước', 'Khác', 'Ly', '=10000+1000'], ['Nước', 'Khác', 'Ly', 'NaN'],
        ]
        for row in invalid:
            with self.subTest(row=row):
                result = self.preview(self.workbook([['Nước', 'Trà', 'Ly', 20000], row])).json()
                self.assertFalse(result['ok'])
                self.assertNotIn('preview_token', result)
                self.assertEqual(result['rows'][1]['action'], 'error')
                self.assertEqual(len(result['rows']), 2)
        self.assertFalse(Product.objects.exists())

    def test_missing_headers_empty_unknown_and_corrupt_files(self):
        for raw in (self.workbook([], None), self.workbook([['Nước', 'Trà']], ['ten_danh_muc', 'ten_san_pham']), b'bad zip'):
            with self.subTest(raw=raw[:10]):
                result = self.preview(raw).json()
                self.assertFalse(result['ok'])
                self.assertNotIn('preview_token', result)
        result = self.preview(self.workbook([['Nước', 'Trà', 'Ly', 100]], ['ten_danh_muc', 'ten_san_pham', 'gia', 'gia'])).json()
        self.assertFalse(result['ok'])
        self.assertIn('trùng tên cột', ' '.join(result['errors']))

    def test_ambiguous_database_names_block_import(self):
        cat, product, _ = self.existing()
        duplicate = Product.objects.create(tenant=self.tenant, category=cat, name='Trà')
        self.assertFalse(self.preview().json()['ok'])
        duplicate.delete()
        Category.objects.create(tenant=self.tenant, name=cat.name)
        self.assertFalse(self.preview().json()['ok'])

    def test_limit_and_no_active_store_block_preview(self):
        self.tenant.max_products = 0
        self.tenant.save()
        result = self.preview().json()
        self.assertFalse(result['ok'])
        self.assertIn('vượt giới hạn', ' '.join(result['errors']))
        self.store.is_active = False
        self.store.save()
        self.assertFalse(self.preview().json()['ok'])

    def test_missing_tampered_expired_and_changed_file_tokens_cannot_write(self):
        raw = self.workbook()
        result = self.preview(raw).json()
        for token in ('', result['preview_token'] + 'tamper'):
            self.assertEqual(self.confirm(raw, token).status_code, 409)
        changed = self.workbook([['Nước', 'Khác', 'Ly', 100]])
        self.assertEqual(self.confirm(changed, result['preview_token']).status_code, 409)
        with patch('django.core.signing.time.time', return_value=9999999999):
            self.assertEqual(self.confirm(raw, result['preview_token']).status_code, 409)
        self.assertFalse(Product.objects.exists())

    def test_catalog_changed_after_scan_requires_rescan(self):
        _, _, unit = self.existing()
        raw = self.workbook()
        result = self.preview(raw).json()
        unit.price = 18000
        unit.save()
        response = self.confirm(raw, result['preview_token'])
        self.assertEqual(response.status_code, 409)
        unit.refresh_from_db()
        self.assertEqual(unit.price, Decimal('18000'))
        result = self.preview(raw).json()
        self.assertEqual(Decimal(result['rows'][0]['old_price']), Decimal('18000'))
        self.assertTrue(self.confirm(raw, result['preview_token']).json()['ok'])

    def test_token_cannot_be_used_by_another_manager_or_tenant(self):
        raw = self.workbook()
        token = self.preview(raw).json()['preview_token']
        other_tenant = Tenant.objects.create(name='Other', public_slug='other-preview')
        other_store = Store.objects.create(tenant=other_tenant, name='Other')
        self.manager.role = User.Role.STAFF
        self.manager.save()
        for tenant, store in ((self.tenant, self.store), (other_tenant, other_store)):
            manager = User.objects.create_user(username=f'manager-{tenant.pk}', tenant=tenant, role=User.Role.MANAGER)
            UserStoreAccess.objects.create(user=manager, store=store, is_default=True)
            self.client.force_login(manager)
            self.assertEqual(self.confirm(raw, token).status_code, 409)
        self.assertFalse(Product.objects.exists())

    def test_post_only_and_permissions(self):
        self.assertEqual(self.client.get(reverse('App_Quanly:catalog_import_preview')).status_code, 405)
        self.manager.role = User.Role.STAFF
        self.manager.save()
        self.assertEqual(self.preview().status_code, 403)
        self.assertEqual(self.confirm(self.workbook()).status_code, 403)

    def test_wrong_extension_oversize_and_missing_file(self):
        for upload in (SimpleUploadedFile('test.csv', b'a,b'), SimpleUploadedFile('test.xlsx', b'x' * (5 * 1024 * 1024 + 1))):
            response = self.client.post(reverse('App_Quanly:catalog_import_preview'), {'excel_file': upload})
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.post(reverse('App_Quanly:catalog_import_preview')).status_code, 400)

    def test_confirmation_rolls_back_if_write_fails(self):
        raw = self.workbook()
        token = self.preview(raw).json()['preview_token']
        with patch('App_Quanly.catalog_excel.ProductUnit.objects.create', side_effect=RuntimeError('write failed')):
            response = self.confirm(raw, token)
        self.assertFalse(response.json()['ok'])
        self.assertFalse(Category.objects.exists())
        self.assertFalse(Product.objects.exists())

    def test_old_topping_workbook_still_previews_and_imports(self):
        self.existing()
        wb = Workbook()
        wb.active.title = 'Topping'
        wb.active.append(['ten_topping', 'hoat_dong'])
        wb.active.append(['Sữa', 1])
        ws = wb.create_sheet('San_pham_Topping')
        ws.append(['ten_danh_muc', 'ten_san_pham', 'ten_topping', 'gia_them', 'hoat_dong'])
        ws.append(['Nước', 'Trà', 'Sữa', 5000, 1])
        buf = io.BytesIO()
        wb.save(buf)
        result = self.preview(buf.getvalue()).json()
        self.assertTrue(result['ok'], result['errors'])
        self.assertTrue(result['warnings'])
        self.assertTrue(self.confirm(buf.getvalue(), result['preview_token']).json()['ok'])
        self.assertEqual(Topping.objects.count(), 1)
        self.assertEqual(ProductTopping.objects.get().price, Decimal('5000'))
