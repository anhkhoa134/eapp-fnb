"""Xuất báo cáo đơn bán ra Excel / CSV (Lịch sử đơn, Dashboard) theo đúng bộ lọc đang chọn."""

import csv
import io
from decimal import Decimal

from django.db.models import Count, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from App_Sales.models import Order, OrderItem

# Giới hạn an toàn: xuất nhiều hơn thì thu hẹp bộ lọc ngày.
MAX_EXPORT_ORDERS = 50000

ORDER_COLUMNS = [
    ('Mã đơn', 16),
    ('Ngày giờ', 18),
    ('Cửa hàng', 22),
    ('Bàn / Loại', 14),
    ('Thu ngân', 18),
    ('Khách hàng', 22),
    ('SĐT khách', 14),
    ('Thanh toán', 12),
    ('Mã giao dịch', 18),
    ('Trạng thái', 14),
    ('Tạm tính (đ)', 14),
    ('Giảm giá (đ)', 14),
    ('Khuyến mãi / ưu đãi', 22),
    ('Thuế (đ)', 12),
    ('Tổng cộng (đ)', 14),
    ('Đã hoàn (đ)', 12),
    ('Thực thu (đ)', 14),
]
MONEY_COLUMNS = {'Tạm tính (đ)', 'Giảm giá (đ)', 'Thuế (đ)', 'Tổng cộng (đ)', 'Đã hoàn (đ)', 'Thực thu (đ)'}


def _int(value) -> int:
    return int((value or Decimal('0')).quantize(Decimal('1')))


def _order_row(order: Order):
    if order.table_name:
        place = order.table_name
    else:
        place = order.get_sale_channel_display() if order.sale_channel else ''
    if order.discount_source == Order.DiscountSource.TIER:
        discount_name = f'Hạng {order.tier_snapshot}'.strip()
    else:
        discount_name = order.promotion_snapshot_name or ''
    refunded = order.refunded_amount or Decimal('0')
    return [
        order.order_code,
        timezone.localtime(order.created_at).strftime('%d/%m/%Y %H:%M'),
        order.store.name,
        place,
        order.cashier.username if order.cashier_id else '',
        order.customer.name if order.customer_id else '',
        order.customer.phone if order.customer_id else '',
        order.get_payment_method_display(),
        order.payment_reference,
        order.get_status_display(),
        _int(order.subtotal),
        _int(order.discount_amount),
        discount_name,
        _int(order.tax_amount),
        _int(order.total_amount),
        _int(refunded),
        _int(order.total_amount - refunded),
    ]


def export_filename(prefix: str, filters: dict, ext: str) -> str:
    parts = [prefix]
    if filters.get('date_from'):
        parts.append(filters['date_from'])
    if filters.get('date_to'):
        parts.append(filters['date_to'])
    if len(parts) == 1:
        parts.append(timezone.localdate().isoformat())
    return '_'.join(parts) + f'.{ext}'


def build_orders_csv(orders) -> bytes:
    buffer = io.StringIO()
    # BOM để Excel trên Windows mở đúng tiếng Việt.
    buffer.write('﻿')
    writer = csv.writer(buffer)
    writer.writerow([title for title, _width in ORDER_COLUMNS])
    for order in orders[:MAX_EXPORT_ORDERS]:
        writer.writerow(_order_row(order))
    return buffer.getvalue().encode('utf-8')


def _style_header(ws, widths):
    fill = PatternFill('solid', fgColor='047857')
    for index, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'A2'


def _money_format(ws, columns, titles, first_row=2):
    for index, title in enumerate(titles, start=1):
        if title in columns:
            for row in ws.iter_rows(min_row=first_row, min_col=index, max_col=index):
                row[0].number_format = '#,##0'


def build_orders_xlsx(orders, *, tenant_name: str, period_label: str) -> bytes:
    order_ids = list(orders.values_list('id', flat=True)[:MAX_EXPORT_ORDERS])
    wb = Workbook()

    ws = wb.active
    ws.title = 'Đơn hàng'
    titles = [title for title, _width in ORDER_COLUMNS]
    ws.append(titles)
    for order in orders[:MAX_EXPORT_ORDERS]:
        ws.append(_order_row(order))
    _style_header(ws, [width for _title, width in ORDER_COLUMNS])
    _money_format(ws, MONEY_COLUMNS, titles)
    last = ws.max_row
    if last > 1:
        ws.append([])
        total_row = ['TỔNG', f'{last - 1} đơn'] + [''] * (len(titles) - 2)
        for index, title in enumerate(titles):
            if title in MONEY_COLUMNS:
                letter = get_column_letter(index + 1)
                total_row[index] = f'=SUM({letter}2:{letter}{last})'
        ws.append(total_row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        _money_format(ws, MONEY_COLUMNS, titles, first_row=ws.max_row)

    # Báo cáo theo món: số lượng + doanh thu dòng món (trước giảm giá đơn).
    by_item = (
        OrderItem.objects.filter(order_id__in=order_ids)
        .values('snapshot_product_name', 'snapshot_unit_name')
        .annotate(quantity=Sum('quantity'), revenue=Sum('line_total'), order_count=Count('order', distinct=True))
        .order_by('-revenue', 'snapshot_product_name')
    )
    ws_items = wb.create_sheet('Theo món')
    item_titles = ['Món', 'Đơn vị', 'Số lượng', 'Số đơn', 'Doanh thu món (đ)']
    ws_items.append(item_titles)
    for row in by_item:
        ws_items.append(
            [
                row['snapshot_product_name'],
                row['snapshot_unit_name'],
                row['quantity'] or 0,
                row['order_count'] or 0,
                _int(row['revenue']),
            ]
        )
    _style_header(ws_items, [30, 14, 12, 10, 18])
    _money_format(ws_items, {'Doanh thu món (đ)'}, item_titles)

    # Theo thu ngân / nhân viên.
    by_cashier = (
        Order.objects.filter(id__in=order_ids)
        .values('cashier__username', 'payment_method')
        .annotate(order_count=Count('id'), total=Sum('total_amount'))
        .order_by('cashier__username', 'payment_method')
    )
    ws_staff = wb.create_sheet('Theo thu ngân')
    staff_titles = ['Thu ngân', 'Thanh toán', 'Số đơn', 'Tổng cộng (đ)']
    ws_staff.append(staff_titles)
    payment_labels = dict(Order.PaymentMethod.choices)
    for row in by_cashier:
        ws_staff.append(
            [
                row['cashier__username'] or '—',
                payment_labels.get(row['payment_method'], row['payment_method']),
                row['order_count'],
                _int(row['total']),
            ]
        )
    _style_header(ws_staff, [22, 14, 10, 16])
    _money_format(ws_staff, {'Tổng cộng (đ)'}, staff_titles)

    wb.properties.title = f'Báo cáo bán hàng {tenant_name} {period_label}'.strip()
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
